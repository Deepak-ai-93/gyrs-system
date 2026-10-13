"""GYRS Job Portal — proper system + UI/UX Word document (job portal only).

Generates docs/GYRS_Portal_Document.docx: scope, sitemap, screen-by-screen
explanations (matching GYRS_Portal_UI_UX.pdf), candidate + admin flows,
add-job field spec, data model, reports spec, roles and deployment.
No tenders section, no OJAS — Apply opens each board's own official website.
"""
import os

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "GYRS_Portal_Document.docx")

NAVY = RGBColor(0, 40, 142)
DARK = RGBColor(15, 23, 42)
GREY = RGBColor(51, 65, 85)

doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hexcolor}"/>'))

def h1(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NAVY; r.font.name = "Arial"; return p

def h2(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = DARK; r.font.name = "Arial"; return p

def para(t, bold_head=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
    if bold_head:
        b = p.add_run(bold_head); b.bold = True; b.font.size = Pt(10); b.font.name = "Arial"; b.font.color.rgb = DARK
    r = p.add_run(t); r.font.size = Pt(10); r.font.name = "Arial"; r.font.color.rgb = GREY; return p

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it); r.font.size = Pt(10); r.font.name = "Arial"; r.font.color.rgb = GREY

def table(headers, rows):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = htxt; shade(c, "00288E")
        for r in c.paragraphs[0].runs:
            r.bold = True; r.font.size = Pt(9.5); r.font.name = "Arial"; r.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri + 1].cells[ci]; c.text = str(val); shade(c, "F8FAFC" if ri % 2 else "FFFFFF")
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(9); r.font.name = "Arial"; r.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


title = doc.add_paragraph(); title.paragraph_format.space_after = Pt(2)
rt = title.add_run("GYRS Job Portal"); rt.bold = True; rt.font.size = Pt(22); rt.font.color.rgb = NAVY; rt.font.name = "Arial"
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(8)
rs = sub.add_run("System Documentation + Screen-by-Screen UI/UX Explanation (companion to GYRS_Portal_UI_UX.pdf)")
rs.font.size = Pt(11); rs.font.color.rgb = RGBColor(100, 116, 139); rs.font.name = "Arial"

table(["Parameter", "Detail"],
      [["Scope", "Job portal only (tender website was a UI/UX reference, not a feature)"],
       ["Apply model", "Candidate reads everything on GYRS; Apply opens the recruiting board's own official website. GYRS never takes applications or fees."],
       ["Third-party portals", "None shown to candidates — no OJAS-style platform anywhere in the UI"],
       ["Stack", "Static HTML + Tailwind CDN + vanilla JS + localStorage (zero-config Vercel deploy)"],
       ["Date / Version", "13 October 2026 — v2 (dashboards + reports)"]])

h1("1. Sitemap — every route and who uses it")
table(["Route", "Surface", "Needs login", "Purpose"],
      [["/", "Home", "No", "Search, live jobs, eligibility calculator, how-it-works"],
       ["/jobs/", "All Jobs", "No", "Filterable listing of every recruitment"],
       ["/job-detail/?id=…", "Job Detail", "No", "Full info + Notification tab + Documents tab + eligibility + official Apply"],
       ["/auth/", "Login / Register", "No", "Candidate + admin entry (demo accounts below)"],
       ["/dashboard/", "User Dashboard", "Yes (user)", "Profile, CV skill-match, saved/applied tracking"],
       ["/skills/", "Skills", "No (saving needs login)", "Upgrade catalog that unlocks jobs"],
       ["/admin/", "Admin Panel", "Yes (admin)", "Filtered dashboard, jobs, users, applications, reports, activity, add-job form"]])

h1("2. Screen-by-screen explanation (as shown in the PDF)")
h2("Screen 1 — Home: the discovery gateway")
para("Hero with trust pill and live vacancy count, 3-part search (keyword, board, qualification), trending pills, Top Recruitment Boards strip, live job cards (vacancies, salary, eligibility, deadline), eligibility calculator, 3-step how-it-works and stats. One screen answers: what jobs exist, am I eligible, how do I apply — every card routes inward to its own Job Detail page.")
h2("Screen 2 — All Jobs: the filterable listing")
para("Keyword + district + sort bar, popular-filter pills, live result counter, sticky sidebar (urgency toggle, board, qualification, category, pay), WhatsApp alerts card, rich cards with metric matrix and syllabus links. The urgency toggle surfaces closing-soon posts first so no deadline is missed.")
h2("Screen 3 — Job Detail: notification, documents, eligibility")
para("Metric grid (vacancies, qualification, salary, location, age, fees, dates), match % for logged-in users, Save / Mark-as-Applied, tabbed official Notification text with PDF download, Documents checklist, and an age/category eligibility checker. The green Apply button opens only that board's official website in a new tab.")
h2("Screen 4 — User Dashboard: profile, CV match, tracking")
para("KPI cards (matched / saved / applied / skills), editable profile, CV upload (.txt) plus paste-text analyser that extracts skills and instantly re-scores every job by skill overlap, CV-ranked recommendations, saved and applied lists. Missing skills point to the Skills page.")
h2("Screen 5 — Skills: upgrade to unlock jobs")
para("Cards (CCC, typing, ITI, nursing, GPSC prep, PET fitness, Tally+GST, Excel), each listing the recruitments it unlocks, with one-click Add to my profile — immediately raising the dashboard match score for those jobs.")
h2("Screen 6 — Admin Dashboard: filters, charts, closing-soon")
para("Global search + board + status + closing-window filters drive KPI cards, vacancies-by-board and jobs-by-status charts, and a closing-soon action list with CSV export. Any slice (e.g. Police + closing in 7 days) can be isolated and downloaded for reminder campaigns.")
h2("Screen 7 — Admin Manage Jobs")
para("Searchable, status-filterable table with views/clicks telemetry, status pills and per-row View / Edit / Duplicate / Status-toggle / Delete, plus jobs CSV export. Full lifecycle draft → published → closed in one table.")
h2("Screen 8 — Admin Add / Edit Job form")
para("All fields: title, board, advt no, posts, qualification, age band, salary, location, fees, start/last/exam dates, skills, notification text, documents checklist, official Apply URL, notification PDF file picker + URL, status. Saving publishes instantly everywhere; every change is written to the Activity Log.")

h1("3. Candidate flow")
bullets(["Discover on Home / All Jobs with full facts visible up front.",
         "Verify on Job Detail: notification tab, documents tab, eligibility checker.",
         "Apply externally on the board's own official website (new tab).",
         "Track on Dashboard (saved/applied) and upskill via Skills to unlock more jobs."])

h1("4. Admin operating guide")
bullets(["Dashboard: filter to any slice, read charts, export the closing-soon list.",
         "Add New Job: fill every field, attach/paste the notification PDF link, Save — live instantly.",
         "Manage Jobs: edit, duplicate repeat recruitments, toggle status, delete, export CSV.",
         "Manage Users: search, filter active/blocked, block/unblock, delete, export CSV.",
         "Applications: every saved + applied event per user/job, filterable, exportable.",
         "Reports: jobs / users / applications / closing-soon CSVs (UTF-8 with BOM for Excel).",
         "Activity Log: timestamped record of all admin actions, clearable."])

h1("5. Add-Job field specification")
table(["Field", "Required", "Example"],
      [["Job title", "Yes", "Staff Nurse Class-III"],
       ["Board / Dept", "Yes", "GSSSB"],
       ["Advt No", "Yes", "GSSSB/202627/192"],
       ["Total posts", "Yes", "1200"],
       ["Qualification", "Yes", "GNM / B.Sc Nursing"],
       ["Age min–max", "No (default 18–35)", "21-38"],
       ["Salary", "No", "₹26,000 Fix Pay"],
       ["Location", "No", "Civil Hospitals Gujarat"],
       ["Fees Gen/Reserved", "No", "500/250"],
       ["Start / Last / Exam dates", "Last recommended", "2026-09-10 / 2026-10-12"],
       ["Skills (comma-separated)", "No — drives CV matching", "patient care, anatomy"],
       ["Notification text", "No", "Official gazette summary…"],
       ["Documents (comma-separated)", "No", "GNM certificate, Photo + signature"],
       ["Official Apply URL", "Yes", "Board's own website (never a third-party portal)"],
       ["Notification PDF", "No", "File picker auto-fills uploads/notifications/….pdf"],
       ["Status", "Yes", "published / draft / closed"]])

h1("6. Data model (browser localStorage)")
table(["Key", "Content"],
      [["gyrs_jobs", "Array of job objects (all fields above + views, clicks)"],
       ["gyrs_users", "Users: name, email, mobile, password, role, education, category, district, skills[], saved[], applied[], blocked"],
       ["gyrs_session", "{email, role} of logged-in user"],
       ["gyrs_activity", "Admin activity log (max 200 entries)"]])

h1("7. Roles and demo accounts")
table(["Role", "Email", "Password", "Lands on"],
      [["Candidate", "priya@example.com", "user123", "/dashboard/"],
       ["Admin", "admin@gyrs.in", "admin123", "/admin/"]])

h1("8. Deploy")
para("Zero-config static hosting. Push to GitHub, import in Vercel with Framework Preset Other, no build command, output root. Auth/persistence is demo-grade localStorage; production hardening (server auth + database) is the defined next step, not part of this UI/UX deliverable.")

doc.save(OUT)
print(f"saved {OUT} ({os.path.getsize(OUT)} bytes)")
